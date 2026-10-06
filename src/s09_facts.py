"""Step 9 — pull the headline figures out of the four recalculated workbooks into
work/facts.json and outputs/Facts_Summary.md, so the report narrative can be
updated against the current numbers.  Run AFTER tools/recalc.py has been applied
to every workbook (run_all.py does this), otherwise formula cells read as None.
"""
import sys, json
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from config import OUTPUTS, WORK, SNAPSHOT_LABEL

def kv(ws, r1, r2, kcol=1, vcol=2, pcol=3):
    out = {}
    for r in range(r1, r2 + 1):
        k = ws.cell(row=r, column=kcol).value
        if k: out[str(k).strip()] = [ws.cell(row=r, column=vcol).value, ws.cell(row=r, column=pcol).value]
    return out
def rules(ws, cols):
    out = {}
    for r in range(5, ws.max_row + 1):
        rid = ws.cell(row=r, column=1).value
        if rid and len(str(rid)) == 3: out[rid] = [ws.cell(row=r, column=c).value for c in cols]
    return out

F = {"snapshot": SNAPSHOT_LABEL}
wb = load_workbook(OUTPUTS["wo"], data_only=True)
F["wo_summary"] = kv(wb["Summary"], 5, 40); F["wo_rules"] = rules(wb["Gap Rules"], (2, 5, 6, 7)); F["backlog"] = kv(wb["Open Backlog"], 5, 26)
wb = load_workbook(OUTPUTS["pm"], data_only=True)
F["pm_summary"] = kv(wb["Summary"], 5, 28); F["pm_rules"] = rules(wb["Gap Rules"], (2, 4, 5, 6))
wb = load_workbook(OUTPUTS["assets"], data_only=True)
F["asset_summary"] = kv(wb["Summary"], 5, 26); F["asset_rules"] = rules(wb["Gap Rules"], (2, 5, 6, 7))
wb = load_workbook(OUTPUTS["inv"], data_only=True)
F["inv_summary"] = kv(wb["Summary"], 5, 21); F["inv_rules"] = rules(wb["Gap Rules"], (2, 4, 5, 6))
try: F["users"] = json.load(open(WORK["user_facts"]))
except FileNotFoundError: pass
json.dump(F, open(WORK["facts"], "w"), indent=1, default=str)

def fmt(v):
    if isinstance(v, float): return f"{v*100:.1f}%" if 0 < abs(v) <= 1 else f"{v:,.1f}"
    if isinstance(v, int): return f"{v:,}"
    return "" if v is None else str(v)
lines = [f"# Facts summary — snapshot {SNAPSHOT_LABEL}", "", "Values read from the recalculated workbooks. Use these to revise the narrative in report/build_report.js.", ""]
for title, key in [("Work orders — Summary", "wo_summary"), ("Open backlog", "backlog"), ("PM schedules — Summary", "pm_summary"), ("Asset register — Summary", "asset_summary"), ("Inventory — Summary", "inv_summary")]:
    lines += [f"## {title}", "", "| Measure | Value | % |", "|---|---|---|"]
    for k, (v, p) in F[key].items():
        if v is not None: lines.append(f"| {k} | {fmt(v)} | {fmt(p) if isinstance(p, float) else ''} |")
    lines.append("")
for title, key, hdr in [("Work order rules", "wo_rules", "Applicable | Fails | Rate"), ("PM rules", "pm_rules", "Applicable | Fails | Rate"), ("Asset rules", "asset_rules", "Applicable | Fails | Rate"), ("Inventory rules", "inv_rules", "Applicable | Fails | Rate")]:
    lines += [f"## {title}", "", f"| Rule | Description | {hdr} |", "|---|---|---|---|---|"]
    for rid, vals in F[key].items():
        lines.append(f"| {rid} | {vals[0]} | {fmt(vals[1])} | {fmt(vals[2])} | {fmt(vals[3])} |")
    lines.append("")
if "users" in F:
    lines += ["## Users (aggregate)", "", "| Measure | Value |", "|---|---|"] + [f"| {k} | {v} |" for k, v in F["users"].items()] + [""]
(OUTPUTS["wo"].parent / "Facts_Summary.md").write_text("\n".join(lines), encoding="utf-8")
print("facts written:", WORK["facts"], "and outputs/Facts_Summary.md")
